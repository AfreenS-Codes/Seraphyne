import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor

def create_project_docx(output_paths):
    doc = Document()

    # Page setup - 1 inch margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    NAVY = RGBColor(10, 25, 47)
    TEAL = RGBColor(13, 148, 136)
    SLATE = RGBColor(71, 85, 105)
    CHARCOAL = RGBColor(30, 41, 59)

    # Document Title
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(0)
    p_pre.paragraph_format.space_after = Pt(4)
    run_pre = p_pre.add_run("SERAPHYNE CLINICAL REASONING PLATFORM")
    run_pre.font.name = "Arial"
    run_pre.font.size = Pt(10)
    run_pre.font.bold = True
    run_pre.font.color.rgb = TEAL

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(8)
    run_title = p_title.add_run("Project Description & Architectural Overview")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(18)
    run_sub = p_sub.add_run("Adaptive Medical Simulation & Cardiovascular Reasoning Engine")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.italic = True
    run_sub.font.color.rgb = SLATE

    # Section 1: Website Description
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    run_h1 = h1.add_run("Website Description")
    run_h1.font.name = "Arial"
    run_h1.font.size = Pt(16)
    run_h1.font.bold = True
    run_h1.font.color.rgb = NAVY

    p1 = doc.add_paragraph()
    p1.paragraph_format.line_spacing = 1.15
    p1.paragraph_format.space_after = Pt(10)
    r1 = p1.add_run(
        "This project is a modern, responsive, and interactive website designed to provide a smooth "
        "and visually engaging user experience. The website combines HTML, CSS, and JavaScript to create "
        "a structured layout, attractive visual styling, interactive components, and smooth scrolling animations."
    )
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.color.rgb = CHARCOAL

    p2 = doc.add_paragraph()
    p2.paragraph_format.line_spacing = 1.15
    p2.paragraph_format.space_after = Pt(10)
    r2 = p2.add_run(
        "The structure of the website is developed using HTML5, with different sections organized to create "
        "a clear and user-friendly navigation flow. CSS3 is used for the overall design, including layouts, "
        "typography, colors, spacing, responsive design, hover effects, transitions, and animated elements. "
        "JavaScript adds interactivity and dynamic behavior to the website, such as scroll-based animations, "
        "navigation interactions, buttons, and other user-triggered effects."
    )
    r2.font.name = "Arial"
    r2.font.size = Pt(11)
    r2.font.color.rgb = CHARCOAL

    p3 = doc.add_paragraph()
    p3.paragraph_format.line_spacing = 1.15
    p3.paragraph_format.space_after = Pt(10)
    r3 = p3.add_run(
        "A key feature of the website is its smooth scrolling and scroll-reveal animations, where elements "
        "progressively appear, move, or transform as the user navigates through the page. These animations are "
        "implemented using a combination of CSS transitions/keyframes and JavaScript scroll detection (Intersection Observer), "
        "creating a more dynamic experience without making the interface unnecessarily complicated."
    )
    r3.font.name = "Arial"
    r3.font.size = Pt(11)
    r3.font.color.rgb = CHARCOAL

    p4 = doc.add_paragraph()
    p4.paragraph_format.line_spacing = 1.15
    p4.paragraph_format.space_after = Pt(16)
    r4 = p4.add_run(
        "The website also focuses on responsive design, allowing the layout and components to adapt cleanly across "
        "desktop, tablet, and mobile screen sizes. Additional design techniques such as Flexbox/Grid, positioning, "
        "gradients, shadows, transitions, transforms, and reusable CSS classes are used to achieve a polished visual appearance."
    )
    r4.font.name = "Arial"
    r4.font.size = Pt(11)
    r4.font.color.rgb = CHARCOAL

    # Section 2: Technologies & Concepts Used
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(8)
    run_h2 = h2.add_run("Technologies & Concepts Used")
    run_h2.font.name = "Arial"
    run_h2.font.size = Pt(16)
    run_h2.font.bold = True
    run_h2.font.color.rgb = NAVY

    techs = [
        ("HTML5", "Website structure and semantic sections"),
        ("CSS3", "Styling, layouts, responsiveness, transitions and animations"),
        ("JavaScript & TypeScript", "Interactivity, type safety, and dynamic functionality"),
        ("CSS Flexbox & Grid", "Flexible and responsive layouts"),
        ("DOM Manipulation", "Controlling webpage elements dynamically"),
        ("Intersection Observer API", "Detecting scrolling and triggering high-performance reveal animations"),
        ("CSS Transitions & Keyframes", "Smooth visual micro-effects and state transforms"),
        ("Responsive Web Design", "Seamless compatibility across desktop, tablet, and mobile screens"),
        ("Hover & Interaction Effects", "Enhanced tactile user engagement with clinical confidence metrics"),
        ("Modern UI Design Principles", "Spacing, typography, hierarchy, and visual consistency (Seraphyne aesthetic)"),
    ]

    for title, desc in techs:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(4)
        run_bold = p_item.add_run(f"{title} – ")
        run_bold.font.name = "Arial"
        run_bold.font.bold = True
        run_bold.font.color.rgb = NAVY
        run_desc = p_item.add_run(desc)
        run_desc.font.name = "Arial"
        run_desc.font.color.rgb = CHARCOAL

    # Section 3: Seraphyne Clinical Features
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(8)
    run_h3 = h3.add_run("Seraphyne Clinical Features & Capabilities")
    run_h3.font.name = "Arial"
    run_h3.font.size = Pt(16)
    run_h3.font.bold = True
    run_h3.font.color.rgb = NAVY

    features = [
        ("Clinical Technology Aesthetic", "White/light-neutral background, deep navy typography, medical teal accents, soft shadows, rounded 12–16px corners, and clean typography hierarchy (Inter/Manrope)."),
        ("Multi-Organ Digital Twin", "Interactive biological schematic with 6 modeled organ systems and real-time state alerts (Stable, Watch, Critical)."),
        ("Side-by-Side Case Comparison", "Dual-patient analytical view comparing vitals deltas, diagnostic probabilities, and outcome trajectories simultaneously."),
        ("Model Confidence Visualizer", "Calibrated Bayesian and machine learning confidence scores with likelihood ratio breakdowns and uncertainty bounds."),
        ("Case History & Audit Replay", "Comprehensive record of analyzed patient sessions with diagnostic accuracy scores and decision review."),
        ("Light & Dark Mode Gravity", "Precision-engineered themes with balanced contrast and reduced clinical eye-strain."),
        ("Ask Seraphyne Clinical Assistant", "Interactive AI assistant equipped with 'Golden Words' highlighting essential clinical pearls and high-yield physiological thresholds."),
        ("Counterfactual Trajectory Explorer", "Analytical bar evaluating 'what-if' decision branches against gold-standard resuscitation pathways."),
    ]

    for title, desc in features:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_after = Pt(4)
        run_bold = p_item.add_run(f"{title}: ")
        run_bold.font.name = "Arial"
        run_bold.font.bold = True
        run_bold.font.color.rgb = TEAL
        run_desc = p_item.add_run(desc)
        run_desc.font.name = "Arial"
        run_desc.font.color.rgb = CHARCOAL

    # Conclusion
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    run_h4 = h4.add_run("Conclusion")
    run_h4.font.name = "Arial"
    run_h4.font.size = Pt(16)
    run_h4.font.bold = True
    run_h4.font.color.rgb = NAVY

    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.line_spacing = 1.15
    p_conc.paragraph_format.space_after = Pt(10)
    r_conc = p_conc.add_run(
        "Overall, the project demonstrates how HTML, CSS, and JavaScript can be combined to transform "
        "a simple static webpage into a responsive, interactive, and visually engaging web experience "
        "adapted to the rigorous demands of clinical reasoning and medical decision-making."
    )
    r_conc.font.name = "Arial"
    r_conc.font.size = Pt(11)
    r_conc.font.color.rgb = CHARCOAL

    for out in output_paths:
        os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)
        doc.save(out)
        print(f"Saved document to {out}")

if __name__ == "__main__":
    targets = [
        "c:/Users/harsh/Downloads/seraphyne/seraphyne_project_description.docx",
        "c:/Users/harsh/Downloads/seraphyne/frontend/public/seraphyne_project_description.docx"
    ]
    create_project_docx(targets)
